"""Stream an offline, escaped HTML index and formula-safe Excel workbook."""

import os
import re
from datetime import datetime
from html import escape
from pathlib import Path
from typing import IO, Any
from urllib.parse import quote

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.cell import WriteOnlyCell  # type: ignore[import-untyped]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]

from officeflow.application.task_archive import ArchiveFile, ArchiveOptions, ArchiveTask
from officeflow.infrastructure.database.task_archive_repository import display

STATUS = {
    "active": "진행",
    "pending": "대기",
    "completed": "완료",
    "canceled": "취소",
    "archived": "보관",
}
PRIORITY = {"normal": "보통", "attention": "관심", "important": "중요", "urgent": "긴급"}
_CONTROLS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class HtmlExcelArchiveWriter:
    def __init__(self) -> None:
        self._html: IO[str] | None = None
        self._book: Any = None
        self._closed = False
        self._task_rows = 0
        self._detail_rows = 0
        self._file_rows = 0

    def start(self, folder: Path, options: ArchiveOptions, created: datetime) -> None:
        self.folder, self.options = folder, options
        self._html = (folder / "업무목록.html.part").open("x", encoding="utf-8")
        self._html.write("""<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>OfficeFlow 업무 보관</title><style>
:root{color-scheme:light}*{box-sizing:border-box}body{margin:0;background:#f5f7fb;color:#172033;font:15px/1.6 "Malgun Gothic",sans-serif}main{max-width:1100px;margin:auto;padding:24px}h1{font-size:24px;margin:0}header{position:sticky;top:0;background:#f5f7fb;padding:12px 0;z-index:1}input,select,button{font:inherit;background:white;color:#172033;border:1px solid #becbe0;border-radius:6px;padding:7px;margin:3px}input{width:min(460px,100%)}details{background:white;border:1px solid #d4ddec;border-radius:8px;margin:8px 0;padding:12px}summary{cursor:pointer;font-weight:600}.status{font-size:13px;color:#405373}pre{white-space:pre-wrap;overflow-wrap:anywhere;font:inherit;margin:8px 0}a{color:#165bc7}section{border-top:1px solid #e1e7f1;margin-top:12px}p{overflow-wrap:anywhere}.tools{display:none}.scripted .tools{display:block}[hidden]{display:none!important}</style>
<main><h1>OfficeFlow 업무 보관</h1>""")
        self._html.write(
            f"<p>범위: {escape(options.description)}<br>생성: {escape(display(created, options.timezone))}</p>"
        )
        self._html.write("""<p>앱 없이 읽는 독립 사본입니다. 내부 자료는 삭제하지 않았습니다. 외부 연결 폴더의 내용, 휴지통·해제 첨부·보기 캐시·자동 이력·중요시점 정보 전체는 포함하지 않습니다. 반복 업무는 한 행이며 저장된 회차만 포함합니다. 파일마다 캡처 시각이 다를 수 있습니다. 검색은 표시 문구와 파일명·메모이며 첨부 본문은 검색하지 않습니다.</p>
<header><div class="tools"><input id="query" type="search" placeholder="제목·내용·일지·파일명·메모 검색" aria-label="업무 검색"><button id="clear">검색 지우기</button><select id="status" aria-label="상태 필터"><option value="">모든 상태</option><option value="active">진행</option><option value="pending">대기</option><option value="completed">완료</option><option value="canceled">취소</option><option value="archived">보관</option></select><span id="count" role="status"></span></div><noscript>검색하려면 JavaScript를 허용하세요. 아래에는 모든 업무의 내용이 펼쳐져 있습니다.</noscript></header>""")
        self._book = Workbook(write_only=True)
        self.summary = self._book.create_sheet("업무목록")
        self.details = self._book.create_sheet("상세내용")
        self.files = self._book.create_sheet("파일목록")
        self.info = self._book.create_sheet("보관안내")
        for sheet in self._book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.column_dimensions["A"].width = 15
            sheet.column_dimensions["B"].width = 34
            sheet.column_dimensions["C"].width = 22
            sheet.column_dimensions["D"].width = 70
        self._append(
            self.summary,
            ["업무ID", "제목", "상태", "중요도", "일정", "등록·수정·완료", "첨부 수", "캡처 시각"],
        )
        self._append(self.details, ["업무ID", "내용 종류", "분할 순번", "내용"])
        self._append(self.files, ["업무ID", "이름", "구분", "메모", "보관 이름", "상대 경로"])
        self._append(self.info, ["항목", "내용"])
        self._append(self.info, ["범위", options.description])
        self._append(self.info, ["생성 시각", display(created, options.timezone)])
        self._append(
            self.info,
            [
                "포함 파일",
                f"현재 첨부={options.attachments}, 보관본={options.versions}, 중요시점 파일={options.snapshot_files}",
            ],
        )
        self._append(
            self.info,
            [
                "주의",
                "읽기용 사본이며 백업/복원 파일이 아닙니다. 폴더 내용·휴지통·해제 첨부·캐시·자동 이력·중요시점 정보 전체 제외. 반복 회차는 저장된 회차만 포함. 각 업무/파일은 다른 시각에 캡처될 수 있음.",
            ],
        )

    @staticmethod
    def _append(sheet: Worksheet, values: list[str | int]) -> None:
        cells = []
        for value in values:
            if isinstance(value, str):
                value = _CONTROLS.sub(lambda match: f"[U+{ord(match.group()):04X}]", value)
                if len(value.encode("utf-16-le")) // 2 > 32767:
                    raise ValueError(
                        "Excel 셀 길이를 초과했습니다. 내용/기간을 나누세요. HTML 준비 자료는 미완료 폴더에 남습니다."
                    )
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"  # '=...', '+...', '@...' must remain literal text.
            cells.append(cell)
        sheet.append(cells)

    def task(
        self, item: ArchiveTask, captured: datetime, files: tuple[tuple[ArchiveFile, str], ...]
    ) -> None:
        assert self._html is not None
        status = STATUS.get(item.status, item.status)
        self._html.write(
            f'<details open data-task="{item.id}" data-status="{escape(item.status, quote=True)}"><summary>{escape(item.title)} <span class="status">· {escape(status)} · 파일 {len(files)}개</span></summary>'
        )
        self._html.write(
            f"<p>업무 #{item.id} · {escape(PRIORITY.get(item.priority, item.priority))}<br>{escape(item.dates)}<br>캡처: {escape(display(captured, self.options.timezone))}</p>"
        )
        entries = (
            ("일정", item.schedule),
            ("설명", item.description),
            ("완료 결과", item.result),
            *item.sections,
        )
        for kind, text in entries:
            if not text:
                continue
            self._html.write(
                f"<section><strong>{escape(kind)}</strong><pre>{escape(text)}</pre></section>"
            )
            # Split long values without losing suffixes to Excel's 32,767 character limit.
            for index, start in enumerate(range(0, len(text), 15000), 1):
                self._detail_rows += 1
                if self._detail_rows >= 1_048_575:
                    raise ValueError("Excel 상세 행 제한입니다. 보관 기간을 나누세요.")
                self._append(self.details, [item.id, kind, index, text[start : start + 15000]])
        if files:
            self._html.write("<section><strong>파일</strong><ul>")
        for file, relative in files:
            text = f"{file.role} · {file.name}" + (f" · {file.label}" if file.label else "")
            self._html.write(
                f'<li><a href="{quote(relative, safe="/")}">{escape(text)}</a><pre>{escape(file.memo)}</pre></li>'
            )
            self._file_rows += 1
            if self._file_rows >= 1_048_575:
                raise ValueError("Excel 파일 행 제한입니다. 보관 기간을 나누세요.")
            self._append(
                self.files, [item.id, file.name, file.role, file.memo[:15000], file.label, relative]
            )
            if len(file.memo) > 15000:
                for index, start in enumerate(range(0, len(file.memo), 15000), 1):
                    self._detail_rows += 1
                    self._append(
                        self.details,
                        [
                            item.id,
                            f"파일 메모: {file.name}",
                            index,
                            file.memo[start : start + 15000],
                        ],
                    )
        if files:
            self._html.write("</ul></section>")
        self._html.write("</details>\n")
        self._task_rows += 1
        self._append(
            self.summary,
            [
                item.id,
                item.title,
                status,
                PRIORITY.get(item.priority, item.priority),
                item.schedule,
                item.dates,
                len(files),
                display(captured, self.options.timezone),
            ],
        )

    def finish(self) -> None:
        assert self._html is not None
        self._html.write("""</main><script>
document.documentElement.classList.add('scripted');
const tasks=Array.from(document.querySelectorAll('details[data-task]'));
const texts=tasks.map(t=>t.textContent.toLocaleLowerCase());
tasks.forEach(t=>t.open=false);
const query=document.getElementById('query'),status=document.getElementById('status');
function filter(){const key=query.value.trim().toLocaleLowerCase();let n=0;tasks.forEach((t,i)=>{t.hidden=(key&&!texts[i].includes(key))||(status.value&&t.dataset.status!==status.value);if(!t.hidden)n++;});document.getElementById('count').textContent=n+' / '+tasks.length+'개 업무';}
query.addEventListener('input',filter);status.addEventListener('change',filter);
document.getElementById('clear').addEventListener('click',()=>{query.value='';filter();query.focus();});filter();
</script></html>""")
        self._html.flush()
        os.fsync(self._html.fileno())
        self._html.close()
        self._html = None
        for sheet, rows, columns in (
            (self.summary, self._task_rows, "H"),
            (self.details, self._detail_rows, "D"),
            (self.files, self._file_rows, "F"),
        ):
            sheet.auto_filter.ref = f"A1:{columns}{rows + 1}"
        temporary = self.folder / "업무목록.xlsx.part"
        self._book.save(temporary)
        # Re-open the complete ZIP and read all sheets; malformed output is not success.
        with temporary.open("r+b") as stream:
            os.fsync(stream.fileno())
            workbook = load_workbook(stream, read_only=True, data_only=False)
            try:
                for sheet in workbook:
                    for _row in sheet.iter_rows():
                        pass
            finally:
                workbook.close()
        (self.folder / "업무목록.html.part").rename(self.folder / "업무목록.html")
        temporary.rename(self.folder / "업무목록.xlsx")

    def close(self) -> None:
        if self._html:
            self._html.close()
            self._html = None
        if self._book is not None and not self._closed:
            self._closed = True
            # Close interrupted write-only generators and delete only their own temporary XML.
            for sheet in self._book.worksheets:
                if not sheet.closed:
                    sheet.close()
                if sheet._writer and Path(sheet._writer.out).exists():
                    sheet._writer.cleanup()
            self._book.close()
