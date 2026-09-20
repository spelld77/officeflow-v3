from __future__ import annotations

import os
import tempfile
from collections.abc import Callable
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import (  # type: ignore[import-untyped]
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.worksheet.table import Table, TableStyleInfo  # type: ignore[import-untyped]

from officeflow.application.exporting import ExportCanceledError
from officeflow.domain.enums import TaskPriority, TaskStatus
from officeflow.domain.task import Task

_HEADERS = (
    "ID",
    "제목",
    "설명",
    "상태",
    "중요도",
    "시작",
    "종료",
    "종일",
    "반복 규칙",
    "고정",
    "완료 요약",
    "첨부",
    "생성일",
    "수정일",
)
_STATUS_LABELS = {
    TaskStatus.ACTIVE: "진행",
    TaskStatus.PENDING: "대기",
    TaskStatus.COMPLETED: "완료",
    TaskStatus.CANCELED: "취소",
    TaskStatus.ARCHIVED: "보관",
}
_PRIORITY_LABELS = {
    TaskPriority.NORMAL: "보통",
    TaskPriority.ATTENTION: "관심",
    TaskPriority.IMPORTANT: "중요",
    TaskPriority.URGENT: "긴급",
}


class ExcelTaskExporter:
    def export(
        self,
        tasks: tuple[Task, ...],
        destination: Path,
        *,
        cancel_requested: Callable[[], bool] | None = None,
    ) -> Path:
        target = destination.resolve()
        if target.suffix.casefold() != ".xlsx":
            raise ValueError("Excel 내보내기 파일은 .xlsx 확장자여야 합니다.")
        target.parent.mkdir(parents=True, exist_ok=True)
        handle, temporary_name = tempfile.mkstemp(
            prefix=f".{target.stem}-", suffix=".part.xlsx", dir=target.parent
        )
        os.close(handle)
        temporary = Path(temporary_name)
        try:
            workbook = self._build_workbook(tasks, cancel_requested)
            workbook.save(temporary)
            workbook.close()
            self._verify(temporary, len(tasks))
            if cancel_requested is not None and cancel_requested():
                raise ExportCanceledError("Excel 내보내기를 취소했습니다.")
            os.replace(temporary, target)
            return target
        finally:
            temporary.unlink(missing_ok=True)

    def _build_workbook(
        self,
        tasks: tuple[Task, ...],
        cancel_requested: Callable[[], bool] | None,
    ) -> Workbook:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "업무"
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.tabColor = "274C77"

        sheet.cell(1, 1, "OfficeFlow 업무 목록")
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(_HEADERS))
        sheet.cell(1, 1).font = Font(name="Arial", size=14, bold=True, color="1F2937")
        sheet.cell(2, 1, f"내보낸 업무: {len(tasks)}개")
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(_HEADERS))
        sheet.cell(2, 1).font = Font(name="Arial", size=10, italic=True, color="64748B")

        header_fill = PatternFill("solid", fgColor="274C77")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        separator = Side(style="thin", color="FFFFFF")
        for column, header in enumerate(_HEADERS, start=1):
            cell = sheet.cell(4, column, header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(right=separator)

        for row, task in enumerate(tasks, start=5):
            if cancel_requested is not None and cancel_requested():
                raise ExportCanceledError("Excel 내보내기를 취소했습니다.")
            values = self._task_values(task)
            for column, item_value in enumerate(values, start=1):
                cell = sheet.cell(row, column, item_value)
                if isinstance(item_value, str):
                    cell.data_type = "s"
                cell.font = Font(name="Arial", size=10, color="1F2937")
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=column in {2, 3, 9, 11},
                )
                if column in {6, 7, 13, 14} and isinstance(item_value, datetime):
                    cell.number_format = (
                        "yyyy-mm-dd" if task.all_day and column in {6, 7} else "yyyy-mm-dd hh:mm"
                    )
            sheet.row_dimensions[row].height = 34

        last_row = max(5, 4 + len(tasks))
        if tasks:
            table = Table(displayName="OfficeFlowTasks", ref=f"A4:N{last_row}")
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            sheet.add_table(table)
        else:
            for column in range(1, len(_HEADERS) + 1):
                sheet.cell(5, column, "")

        widths = (9, 30, 42, 10, 10, 19, 19, 9, 30, 9, 42, 9, 19, 19)
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width
        sheet.row_dimensions[1].height = 24
        sheet.row_dimensions[4].height = 24
        sheet.freeze_panes = "A5"
        sheet.auto_filter.ref = f"A4:N{last_row}"
        sheet.print_title_rows = "4:4"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        return workbook

    @staticmethod
    def _task_values(task: Task) -> tuple[object, ...]:
        zone = ZoneInfo(task.timezone)

        def local(value: datetime | None) -> datetime | None:
            return value.astimezone(zone).replace(tzinfo=None) if value is not None else None

        return (
            task.id,
            task.title,
            task.description,
            _STATUS_LABELS[task.status],
            _PRIORITY_LABELS[task.priority],
            local(task.starts_at),
            local(task.ends_at),
            "예" if task.all_day else "아니오",
            task.recurrence_rule or "",
            "예" if task.is_pinned else "아니오",
            task.result_note,
            "있음" if task.has_attachments else "없음",
            local(task.created_at),
            local(task.updated_at),
        )

    @staticmethod
    def _verify(path: Path, expected_rows: int) -> None:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            if workbook.sheetnames != ["업무"]:
                raise OSError("Excel 파일의 시트 구성이 올바르지 않습니다.")
            sheet = workbook["업무"]
            actual_rows = max(0, sheet.max_row - 4)
            if expected_rows == 0:
                actual_rows = 0
            if actual_rows != expected_rows:
                raise OSError("Excel 파일의 업무 개수 검증에 실패했습니다.")
            actual_headers = tuple(sheet.cell(4, index).value for index in range(1, 15))
            if actual_headers != _HEADERS:
                raise OSError("Excel 파일의 열 구성이 올바르지 않습니다.")
        finally:
            workbook.close()
