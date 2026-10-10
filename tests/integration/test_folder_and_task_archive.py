import hashlib
import json
import sqlite3
from dataclasses import replace
from datetime import UTC, date, datetime
from pathlib import Path

import pytest
from alembic import command
from openpyxl import load_workbook
from sqlalchemy import delete, select

from officeflow.application.attachment_export import AttachmentExportService
from officeflow.application.folder_links import FolderLinkService
from officeflow.application.task_archive import (
    ArchiveBundleError,
    ArchiveDateBasis,
    ArchiveOptions,
    ArchiveScope,
    TaskArchiveService,
)
from officeflow.application.tasks import TaskDraft
from officeflow.bootstrap.paths import AppPaths
from officeflow.infrastructure.attachments.export_storage import ExternalExportStorage
from officeflow.infrastructure.backup import BackupManager
from officeflow.infrastructure.database.folder_link_repository import SqlAlchemyFolderLinkRepository
from officeflow.infrastructure.database.migrate import migration_config, upgrade_database
from officeflow.infrastructure.database.models import TaskFolderLinkRecord, TaskRecord
from officeflow.infrastructure.database.task_archive_repository import SqlAlchemyArchiveRepository
from officeflow.infrastructure.exports.task_archive import HtmlExcelArchiveWriter
from tests.integration.test_attachment_documents import documents, seed


def services(env):
    links = FolderLinkService(
        SqlAlchemyFolderLinkRepository(env.sessions),
        (env.database.parent, env.attachments.storage_root),
    )
    exports = AttachmentExportService(env.attachments, ExternalExportStorage(env.database.parent))
    archives = TaskArchiveService(
        SqlAlchemyArchiveRepository(env.sessions), exports, HtmlExcelArchiveWriter
    )
    return links, archives


def selected(identity, **kwargs):
    return ArchiveOptions(
        date(2026, 1, 1),
        date(2026, 12, 31),
        scope=ArchiveScope.SELECTED,
        task_ids=(identity,),
        **kwargs,
    )


def test_folder_metadata_relink_rename_missing_internal_and_cascade(tmp_path):
    with documents(tmp_path / "profile") as (env, _):
        source = seed(env)
        links, _ = services(env)
        outside = tmp_path / "outside"
        outside.mkdir()
        marker = outside / "user.txt"
        marker.write_text("keep", encoding="utf-8")
        link = links.save(source.task_id, "提出", outside)
        assert links.open_path(link) == outside.resolve()
        with pytest.raises(ValueError, match="이미"):
            links.save(source.task_id, "dup", outside)
        with pytest.raises(ValueError, match="관리"):
            links.save(source.task_id, "bad", env.database.parent)
        with pytest.raises(ValueError, match="절대"):
            links.save(source.task_id, "bad", Path("relative"))
        renamed = links.rename(source.task_id, link, "새 이름")
        assert renamed.path == link.path and renamed.id == link.id
        relocated = tmp_path / "relocated"
        outside.rename(relocated)
        with pytest.raises(ValueError, match="연결은 유지"):
            links.open_path(renamed)
        assert links.list(source.task_id)[0].id == link.id
        updated = links.save(source.task_id, renamed.name, relocated, identity=link.id)
        assert updated.position == link.position and updated.created_at == link.created_at
        links.delete(source.task_id, link.id)
        assert (relocated / "user.txt").read_text(encoding="utf-8") == "keep"
        links.save(source.task_id, "또 연결", relocated)
        with env.sessions.transaction() as session:
            # Delete links by FK only, never by traversing user folders.
            session.execute(
                delete(TaskFolderLinkRecord).where(TaskFolderLinkRecord.task_id == source.task_id)
            )
        assert links.list(source.task_id) == () and (relocated / "user.txt").exists()


def test_folder_trash_read_only_and_backup_contains_metadata_not_external_file(tmp_path):
    import zipfile

    paths = AppPaths(tmp_path / "profile")
    paths.ensure_directories()
    with documents(paths.data_dir, attachment_root=paths.attachment_dir) as (env, _):
        source = seed(env)
        links, _ = services(env)
        outside = tmp_path / "company"
        outside.mkdir()
        (outside / "not-backed-up.txt").write_text("private-external", encoding="utf-8")
        link = links.save(source.task_id, "회사 폴더", outside)
        with env.sessions.transaction() as session:
            session.get(TaskRecord, source.task_id).deleted_at = datetime.now(UTC)
        assert links.list(source.task_id)[0] == link
        for operation in (
            lambda: links.rename(source.task_id, link, "new"),
            lambda: links.delete(source.task_id, link.id),
        ):
            with pytest.raises(ValueError, match="휴지통"):
                operation()
        manager = BackupManager(paths)
        backup = manager.create_backup()
        with zipfile.ZipFile(backup.path) as bundle:
            assert not any("not-backed-up" in name for name in bundle.namelist())
            database_member = next(
                name for name in bundle.namelist() if name.endswith("officeflow.db")
            )
            target = tmp_path / "captured.db"
            target.write_bytes(bundle.read(database_member))
        with sqlite3.connect(target) as db:
            assert db.execute("SELECT name,path FROM task_folder_links").fetchone() == (
                link.name,
                link.path,
            )


def test_0015_upgrade_keeps_data_and_adds_links(tmp_path):
    database = tmp_path / "old.db"
    command.upgrade(migration_config(database), "0015_task_snapshots")
    with sqlite3.connect(database) as db:
        assert (
            db.execute("SELECT name FROM sqlite_master WHERE name='task_folder_links'").fetchone()
            is None
        )
        db.execute("INSERT INTO app_settings VALUES ('marker','{}','2026-01-01')")
    upgrade_database(database)
    with sqlite3.connect(database) as db:
        assert (
            db.execute("SELECT version_num FROM alembic_version").fetchone()[0]
            == "0016_task_folder_links"
        )
        assert db.execute("SELECT key FROM app_settings").fetchone()[0] == "marker"
        assert db.execute("PRAGMA foreign_key_check").fetchall() == []
    assert list((tmp_path / "schema-backups").glob("*.db"))


def test_archive_offline_escape_long_text_formula_files_and_no_internal_changes(tmp_path):
    with documents(tmp_path / "profile") as (env, _):
        source = seed(env, b"original")
        links, archives = services(env)
        long_text = "重要" * 25000 + "tail-marker"
        env.tasks.update(
            source.task_id, TaskDraft("=1+1", description="<script>alert(1)</script>\n" + long_text)
        )
        env.records.add_work_log(
            task_id=source.task_id, log_date=date(2026, 1, 7), content="찾을 단어", result="결과"
        )
        env.records.add_checklist_item(source.task_id, "체크 검색")
        output = tmp_path / "archive"
        output.mkdir()
        links.save(source.task_id, "외부 폴더", output)
        path = env.attachments.storage_root / source.relative_path
        path.write_bytes(b"saved current content")
        before = env.attachments.get_metadata(source.id)
        result = archives.export(output, selected(source.task_id))
        assert result.task_count == 1 and result.file_count == 1
        assert (result.folder / "완료.txt").exists() and not (result.folder / "미완료.txt").exists()
        html = (result.folder / "업무목록.html").read_text(encoding="utf-8")
        assert (
            "<script>alert(1)</script>" not in html
            and "&lt;script&gt;alert(1)&lt;/script&gt;" in html
        )
        assert "찾을 단어" in html and "tail-marker" in html and "외부 폴더" in html
        assert "t.textContent.toLocaleLowerCase()" in html
        assert "https://" not in html and "color-scheme:light" in html
        manifest = json.loads((result.folder / "파일목록.jsonl").read_text(encoding="utf-8"))
        saved = result.folder / manifest["path"]
        assert saved.read_bytes() == b"saved current content"
        assert hashlib.sha256(saved.read_bytes()).hexdigest() == manifest["sha256"]
        assert env.attachments.get_metadata(source.id) == before
        book = load_workbook(result.folder / "업무목록.xlsx", read_only=True, data_only=False)
        try:
            assert (
                book["업무목록"]["B2"].value == "=1+1" and book["업무목록"]["B2"].data_type == "s"
            )
            text = "".join(
                row[3]
                for row in book["상세내용"].iter_rows(min_row=2, values_only=True)
                if row[1] == "설명"
            )
            assert text == "<script>alert(1)</script>\n" + long_text
        finally:
            book.close()
        again = archives.export(output, selected(source.task_id, attachments=False))
        assert again.folder != result.folder and again.file_count == 0
        assert path.read_bytes() == b"saved current content"


@pytest.mark.parametrize("basis", list(ArchiveDateBasis))
def test_period_boundaries_unscheduled_completed_filters_and_selected_ignore_dates(tmp_path, basis):
    with documents(tmp_path) as (env, _):
        a = env.tasks.create(TaskDraft("one"))
        b = env.tasks.create(TaskDraft("two"))
        with env.sessions.transaction() as session:
            for task, timestamp in (
                (a, datetime(2025, 12, 31, 15, tzinfo=UTC)),
                (b, datetime(2026, 12, 31, 15, tzinfo=UTC)),
            ):
                row = session.get(TaskRecord, task.id)
                setattr(row, basis.value, timestamp)
                row.status = "completed"
        repo = SqlAlchemyArchiveRepository(env.sessions)
        options = ArchiveOptions(date(2026, 1, 1), date(2026, 12, 31), basis=basis)
        assert repo.task_ids(options) == (a.id,)
        assert repo.task_ids(selected(b.id)) == (b.id,)
        with env.sessions.transaction() as session:
            session.get(TaskRecord, a.id).status = "active"
        assert repo.task_ids(options) == ()
        assert repo.task_ids(replace(options, completed_only=False)) == (a.id,)


@pytest.mark.parametrize("failure", ["missing", "cancel", "change", "write"])
def test_failure_never_marks_complete_keeps_sources_and_releases_pins(
    tmp_path, monkeypatch, failure
):
    with documents(tmp_path / "profile") as (env, _):
        source = seed(env)
        _, archives = services(env)
        output = tmp_path / "archive"
        output.mkdir()
        path = env.attachments.storage_root / source.relative_path
        if failure == "missing":
            path.unlink()
        original = archives.files.storage.prepare
        cancel = [False]

        def prepare(*args, **kwargs):
            prepared = original(*args, **kwargs)
            if failure == "cancel":
                cancel[0] = True
            if failure == "change":
                env.tasks.update(source.task_id, TaskDraft("changed"))
            if failure == "write":
                raise OSError("destination failure")
            return prepared

        monkeypatch.setattr(archives.files.storage, "prepare", prepare)
        with pytest.raises(ArchiveBundleError, match="미완료 폴더"):
            archives.export(output, selected(source.task_id), canceled=lambda: cancel[0])
        bundle = next(output.iterdir())
        assert (bundle / "미완료.txt").exists() and not (bundle / "완료.txt").exists()
        assert not env.attachments.coordinator.is_pinned(source.relative_path)
        assert env.attachments.get_metadata(source.id).id == source.id
        if failure != "missing":
            assert path.read_bytes() == b"initial"


def test_internal_destination_and_deleted_selection_rejected_before_bundle(tmp_path):
    with documents(tmp_path) as (env, _):
        source = seed(env)
        _, archives = services(env)
        with pytest.raises(ValueError, match="관리"):
            archives.export(env.database.parent, selected(source.task_id))
        with env.sessions.transaction() as session:
            session.get(TaskRecord, source.task_id).deleted_at = datetime.now(UTC)
        with pytest.raises(ValueError, match="삭제/휴지통"):
            archives.repository.task_ids(selected(source.task_id))


def test_file_categories_and_fixed_integrity_failure(tmp_path):
    from officeflow.application.attachment_documents import DocumentOperation
    from officeflow.application.task_snapshots import SnapshotCommand, SnapshotSelection
    from tests.integration.test_document_lifecycle import lifecycle, original

    with lifecycle(tmp_path / "profile") as (env, facade):
        item = original(env, facade)
        work = facade.create(
            facade.catalog.get("attachment", item.id), DocumentOperation.WORK, filename="작업.txt"
        )
        facade.create(
            facade.catalog.get("attachment", work.attachment_id),
            DocumentOperation.VERSION,
            label="제출 전",
            memo="보관 메모",
        )
        from uuid import uuid4

        facade.snapshots.create(
            SnapshotCommand(
                uuid4().hex,
                item.task_id,
                "중요시점",
                (SnapshotSelection("attachment", item.id),),
                documents_closed=True,
            )
        )
        _, archives = services(env)
        outside = tmp_path / "exports"
        outside.mkdir()
        result = archives.export(
            outside, selected(item.task_id, versions=True, snapshot_files=True)
        )
        assert result.file_count == 4
        records = [
            json.loads(line)
            for line in (result.folder / "파일목록.jsonl").read_text(encoding="utf-8").splitlines()
        ]
        assert {row["kind"] for row in records} == {"attachment", "version", "snapshot_file"}
        assert len({row["path"] for row in records}) == 4
        for row in records:
            assert (result.folder / row["path"]).read_bytes() == b"initial"
        source = env.attachments.storage_root / item.relative_path
        source.chmod(0o666)
        source.write_bytes(b"corrupted frozen source")
        with pytest.raises(ArchiveBundleError, match="고정 파일"):
            archives.export(outside, selected(item.task_id))
        assert (
            source.read_bytes() == b"corrupted frozen source"
        )  # No automatic repair or destruction.


def test_folder_stale_name_edit_does_not_revert_new_location(tmp_path):
    with documents(tmp_path / "profile") as (env, _):
        task = env.tasks.create(TaskDraft("folder"))
        links, _ = services(env)
        first, second = tmp_path / "first", tmp_path / "second"
        first.mkdir()
        second.mkdir()
        link = links.save(task.id, "name", first)
        moved = links.save(task.id, "name", second, identity=link.id, expected=link)
        with pytest.raises(ValueError, match="다른 창"):
            links.rename(task.id, link, "stale")
        assert links.list(task.id)[0] == moved


def test_folder_backup_restores_metadata_and_task_delete_only_cascades_link(tmp_path):
    from sqlalchemy import delete

    paths = AppPaths(tmp_path / "source")
    paths.ensure_directories()
    with documents(paths.data_dir, attachment_root=paths.attachment_dir) as (env, _):
        task = env.tasks.create(TaskDraft("no files"))
        links, _ = services(env)
        outside = tmp_path / "outside"
        outside.mkdir()
        marker = outside / "marker.txt"
        marker.write_text("must stay", encoding="utf-8")
        link = links.save(task.id, "outside", outside)
        backup = BackupManager(paths).create_backup()
        with env.sessions.transaction() as session:
            session.execute(delete(TaskRecord).where(TaskRecord.id == task.id))
        assert links.list(task.id) == () and marker.exists()
    destination = AppPaths(tmp_path / "restored")
    destination.ensure_directories()
    upgrade_database(destination.database_file)
    manager = BackupManager(destination)
    manager.stage_restore(backup.path)
    manager.apply_pending_restore()
    with sqlite3.connect(destination.database_file) as db:
        assert db.execute("SELECT name,path FROM task_folder_links").fetchone() == (
            link.name,
            link.path,
        )
        assert db.execute("PRAGMA integrity_check").fetchone()[0] == "ok"


def test_repeated_task_is_one_row_with_saved_occurrence_result(tmp_path):
    with documents(tmp_path / "profile") as (env, _):
        start = datetime(2026, 1, 2, 1, tzinfo=UTC)
        task = env.tasks.create(TaskDraft("repeat", starts_at=start, recurrence_rule="FREQ=DAILY"))
        env.tasks.ensure_occurrence(task.id, start)
        from officeflow.infrastructure.database.models import TaskOccurrenceRecord

        with env.sessions.transaction() as session:
            occurrence = session.scalar(
                select(TaskOccurrenceRecord).where(TaskOccurrenceRecord.task_id == task.id)
            )
            occurrence.result_note = "회차 결과"
            occurrence.status = "completed"
        _, archives = services(env)
        item = archives.repository.capture(task.id, selected(task.id))
        assert len([kind for kind, _text in item.sections if kind == "저장된 반복 회차"]) == 1
        assert any("회차 결과" in text for _kind, text in item.sections)


def test_link_and_archive_reject_junction_into_managed_root(tmp_path):
    import subprocess

    with documents(tmp_path / "profile") as (env, _):
        source = seed(env)
        links, archives = services(env)
        alias = tmp_path / "managed-alias"
        result = subprocess.run(
            ["cmd", "/c", "mklink", "/J", str(alias), str(env.database.parent)],
            capture_output=True,
            timeout=10,
        )
        assert result.returncode == 0
        with pytest.raises(ValueError, match="관리"):
            links.save(source.task_id, "alias", alias)
        with pytest.raises(ValueError, match="관리"):
            archives.export(alias, selected(source.task_id))


def test_new_schema_backup_missing_folder_table_is_rejected(tmp_path):
    from officeflow.infrastructure.backup import BackupError, _check_restore_schema

    database = tmp_path / "corrupt.db"
    upgrade_database(database)
    with sqlite3.connect(database) as connection:
        connection.execute("DROP TABLE task_folder_links")
    with pytest.raises(BackupError, match="폴더 연결"):
        _check_restore_schema(database)


@pytest.mark.parametrize("stage", ["copying", "prepared"])
def test_killed_export_stays_incomplete_and_restart_preserves_sources(tmp_path, stage):
    import os
    import subprocess
    import sys

    temp = tmp_path / "temporary"
    temp.mkdir()
    script = """
import os,sys
from pathlib import Path
from tests.integration.test_attachment_documents import documents,seed
from tests.integration.test_folder_and_task_archive import services,selected
root=Path(sys.argv[1]); output=root/'archive';output.mkdir()
with documents(root/'profile') as (env,_):
    source=seed(env)
    _,archives=services(env)
    archives.files.storage._checkpoint=lambda phase,prepared: os._exit(71) if phase==sys.argv[2] else None
    archives.export(output,selected(source.task_id))
"""
    result = subprocess.run(
        [sys.executable, "-c", script, str(tmp_path), stage],
        env={**os.environ, "TEMP": str(temp), "TMP": str(temp)},
        capture_output=True,
        timeout=30,
    )
    assert result.returncode == 71, result.stderr.decode(errors="replace")
    old_bundle = next((tmp_path / "archive").iterdir())
    assert (old_bundle / "미완료.txt").exists() and not (old_bundle / "완료.txt").exists()
    with documents(tmp_path / "profile") as (env, _):
        with env.sessions.transaction() as session:
            identity = session.scalar(select(TaskRecord.id))
        _, archives = services(env)
        item = archives.repository.capture(identity, selected(identity))
        source = archives.files.resolve(item.files[0].reference).path
        assert source.read_bytes() == b"initial"
        complete = archives.export(tmp_path / "archive", selected(identity))
        assert complete.folder != old_bundle and (complete.folder / "완료.txt").exists()
        assert (old_bundle / "미완료.txt").exists()


def test_large_file_archive_memory_does_not_scale_with_file_size(tmp_path):
    import tracemalloc
    from time import perf_counter

    with documents(tmp_path / "profile") as (env, _):
        task = env.tasks.create(TaskDraft("large"))
        file = tmp_path / "large.bin"
        with file.open("wb") as stream:
            for _ in range(32):
                stream.write(b"x" * (1024 * 1024))
        attachment = env.attachments.attach(task.id, file)
        _, archives = services(env)
        output = tmp_path / "archive"
        output.mkdir()
        tracemalloc.start()
        started = perf_counter()
        try:
            result = archives.export(output, selected(task.id))
            _, peak = tracemalloc.get_traced_memory()
        finally:
            tracemalloc.stop()
        print(
            f"archive 32MiB: peak_python={peak / 1024**2:.2f}MiB elapsed={perf_counter() - started:.2f}s"
        )
        assert peak < 24 * 1024**2
        metadata = json.loads((result.folder / "파일목록.jsonl").read_text(encoding="utf-8"))
        assert (result.folder / metadata["path"]).stat().st_size == 32 * 1024**2
        assert not env.attachments.coordinator.is_pinned(attachment.relative_path)


@pytest.mark.parametrize("stage", ["marker_write", "old_marker_cleanup"])
def test_marker_write_failure_cannot_look_complete_and_cleanup_warning_is_not_data_failure(
    tmp_path, monkeypatch, stage
):
    with documents(tmp_path / "profile") as (env, _):
        source = seed(env)
        _, archives = services(env)
        outside = tmp_path / "archive"
        outside.mkdir()
        if stage == "marker_write":
            write = archives._write

            def failed_write(path, text):
                write(path, text)
                if path.name == ".완료.txt.part":
                    raise OSError("simulated marker fsync failure")

            monkeypatch.setattr(archives, "_write", failed_write)
            with pytest.raises(ArchiveBundleError, match="미완료 폴더"):
                archives.export(outside, selected(source.task_id))
            folder = next(outside.iterdir())
            assert (folder / "미완료.txt").exists() and not (folder / "완료.txt").exists()
        else:
            unlink = Path.unlink

            def failed_unlink(path, *args, **kwargs):
                if path.name == "미완료.txt":
                    raise PermissionError("marker held by scanner")
                return unlink(path, *args, **kwargs)

            monkeypatch.setattr(Path, "unlink", failed_unlink)
            result = archives.export(outside, selected(source.task_id))
            assert result.warning and (result.folder / "완료.txt").exists()
