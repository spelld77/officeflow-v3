import tracemalloc
from dataclasses import replace
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from officeflow.application.exporting import ExportCanceledError, ExportService
from officeflow.application.tasks import TaskDraft, TaskQuery, TaskService, TaskSort
from officeflow.domain.task import Task
from officeflow.infrastructure.database.migrate import upgrade_database
from officeflow.infrastructure.database.session import SessionFactory, create_database_engine
from officeflow.infrastructure.database.task_repository import SqlAlchemyTaskRepository
from officeflow.infrastructure.exports.calendar import ICalendarTaskExporter
from officeflow.infrastructure.exports.excel import ExcelTaskExporter


def test_export_pages_read_one_snapshot_while_users_edit_data(tmp_path):
    database = tmp_path / "data.db"
    upgrade_database(database)
    engine = create_database_engine(database)
    sessions = SessionFactory(engine)
    repository = SqlAlchemyTaskRepository(sessions)
    service = TaskService(repository)
    start = datetime(2026, 10, 1, tzinfo=UTC)
    with sessions.transaction() as session:
        session.add_all(
            [
                repository._to_record(
                    Task.create(
                        title=f"업무 {number:04}",
                        description="원래 설명",
                        starts_at=start + timedelta(minutes=number),
                        now=start,
                    )
                )
                for number in range(1205)
            ]
        )
    changed = False

    def progress(_count):
        nonlocal changed
        if not changed:
            changed = True
            service.update(1100, TaskDraft(title="내보내는 도중 수정"))
            service.move_to_trash(1200)
            service.create(TaskDraft(title="내보내는 도중 추가"))

    path = tmp_path / "snapshot.xlsx"
    try:
        ExportService(service, ExcelTaskExporter(), ICalendarTaskExporter()).export_excel(
            TaskQuery(sort=TaskSort.TITLE), path, progress=progress
        )
        book = load_workbook(path, read_only=True)
        try:
            rows = list(book["업무"].iter_rows(min_row=5, values_only=True))
        finally:
            book.close()
        assert len(rows) == 1205
        assert len({row[0] for row in rows}) == 1205
        assert next(row for row in rows if row[0] == 1100)[1] == "업무 1099"
        assert any(row[0] == 1200 for row in rows)
        assert not any(row[1] == "내보내는 도중 추가" for row in rows)
    finally:
        engine.dispose()


def test_streaming_writer_has_bounded_memory_for_large_excel(tmp_path):
    start = datetime(2026, 10, 1, tzinfo=UTC)
    base = Task.create(title="대량 검증", description="긴 설명 " * 200, starts_at=start, now=start)
    counts = []
    tracemalloc.start()
    try:
        ExcelTaskExporter().export_stream(
            (replace(base, id=number, title=f"업무 {number}") for number in range(10_000)),
            tmp_path / "large.xlsx",
            progress=counts.append,
        )
        _current, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert counts[-1] == 10_000
    assert peak < 32 * 1024 * 1024
    print(f"10,000-row streamed Excel Python peak: {peak / 1024**2:.2f} MiB")


@pytest.mark.parametrize("phase", ["write", "verify", "save"])
def test_stream_cancellation_or_failure_preserves_target_and_cleans_temp(
    tmp_path, monkeypatch, phase
):
    import openpyxl.worksheet._writer as writer_module

    start = datetime(2026, 10, 1, tzinfo=UTC)
    task = Task.create(title="안전한 저장", starts_at=start, now=start)
    path = tmp_path / "existing.xlsx"
    path.write_bytes(b"original")
    before = set(writer_module.ALL_TEMP_FILES)
    exporter = ExcelTaskExporter()
    if phase == "save":
        from openpyxl import Workbook

        def fail(*_args, **_kwargs):
            raise OSError("저장 실패")

        monkeypatch.setattr(Workbook, "save", fail)
        with pytest.raises(OSError):
            exporter.export_stream(iter((task,)), path)
    elif phase == "verify":

        def cancel_verify(*_args, **_kwargs):
            raise ExportCanceledError("검증 취소")

        monkeypatch.setattr(exporter, "_verify", cancel_verify)
        with pytest.raises(ExportCanceledError):
            exporter.export_stream(iter((task,)), path)
    else:
        with pytest.raises(ExportCanceledError):
            exporter.export_stream(iter((task,)), path, cancel_requested=lambda: True)
    assert path.read_bytes() == b"original"
    assert not list(tmp_path.glob("*.part.xlsx"))
    assert set(writer_module.ALL_TEMP_FILES) == before
    assert all(Path(name).exists() for name in before)


def test_empty_stream_still_produces_headers_and_valid_workbook(tmp_path):
    path = ExcelTaskExporter().export_stream(iter(()), tmp_path / "empty.xlsx")
    book = load_workbook(path, read_only=True)
    try:
        rows = list(book["업무"].iter_rows(values_only=True))
        assert rows[3][0] == "ID" and rows[3][1] == "제목"
        assert len(rows) == 4
    finally:
        book.close()
