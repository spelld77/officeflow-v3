from __future__ import annotations

from dataclasses import replace
from datetime import UTC, datetime, timedelta

from openpyxl import load_workbook  # type: ignore[import-untyped]

from officeflow.application.exporting import ExportService
from officeflow.application.tasks import TaskPage, TaskQuery
from officeflow.domain.enums import TaskPriority
from officeflow.domain.task import Task
from officeflow.infrastructure.exports.calendar import ICalendarTaskExporter
from officeflow.infrastructure.exports.excel import ExcelTaskExporter


def _task(
    *,
    task_id: int,
    title: str,
    all_day: bool,
    starts_at: datetime,
    ends_at: datetime,
    recurrence_rule: str | None = None,
) -> Task:
    created = datetime(2026, 9, 1, tzinfo=UTC)
    return replace(
        Task.create(
        title=title,
        description="설명, 줄바꿈\n둘째 줄",
        priority=TaskPriority.IMPORTANT,
        all_day=all_day,
        starts_at=starts_at,
        ends_at=ends_at,
        timezone="Asia/Seoul",
        recurrence_rule=recurrence_rule,
        now=created,
        ),
        id=task_id,
    )


def test_excel_export_preserves_typed_schedule_and_blocks_formula_injection(tmp_path) -> None:
    start = datetime(2026, 9, 15, 0, 0, tzinfo=UTC)
    task = _task(
        task_id=7,
        title="=HYPERLINK(\"https://invalid\")",
        all_day=False,
        starts_at=start,
        ends_at=start + timedelta(hours=2),
    )
    destination = tmp_path / "tasks.xlsx"

    ExcelTaskExporter().export((task,), destination)

    workbook = load_workbook(destination, data_only=False)
    try:
        sheet = workbook["업무"]
        assert sheet["B5"].value == task.title
        assert sheet["B5"].data_type == "s"
        assert isinstance(sheet["F5"].value, datetime)
        assert sheet.freeze_panes == "A5"
        assert "OfficeFlowTasks" in sheet.tables
        assert sheet.sheet_view.showGridLines is False
    finally:
        workbook.close()


def test_ics_export_preserves_multiday_end_and_recurrence(tmp_path) -> None:
    start = datetime(2026, 9, 14, 15, 0, tzinfo=UTC)
    multiday = _task(
        task_id=1,
        title="여러 날 일정",
        all_day=True,
        starts_at=start,
        ends_at=start + timedelta(days=3),
    )
    recurring = _task(
        task_id=2,
        title="매주 회의",
        all_day=False,
        starts_at=start,
        ends_at=start + timedelta(hours=1),
        recurrence_rule="FREQ=WEEKLY;INTERVAL=1;BYDAY=TU",
    )
    destination = tmp_path / "schedule.ics"

    ICalendarTaskExporter().export((multiday, recurring), destination)
    payload = destination.read_bytes()

    assert b"\r\n" in payload
    text = payload.decode("utf-8")
    assert "DTSTART;VALUE=DATE:20260915" in text
    assert "DTEND;VALUE=DATE:20260918" in text
    assert "DTSTART;TZID=Asia/Seoul:20260915T000000" in text
    assert "DTEND;TZID=Asia/Seoul:20260915T010000" in text
    assert "RRULE:FREQ=WEEKLY;INTERVAL=1;BYDAY=TU" in text
    assert text.count("BEGIN:VEVENT") == 2


def test_export_cancellation_leaves_no_partial_file(tmp_path) -> None:
    start = datetime(2026, 9, 15, tzinfo=UTC)
    task = _task(
        task_id=1,
        title="취소 확인",
        all_day=True,
        starts_at=start,
        ends_at=start + timedelta(days=1),
    )
    destination = tmp_path / "canceled.xlsx"

    try:
        ExcelTaskExporter().export((task,), destination, cancel_requested=lambda: True)
    except RuntimeError:
        pass
    else:
        raise AssertionError("취소된 내보내기가 성공으로 처리되었습니다.")

    assert not destination.exists()
    assert not tuple(tmp_path.glob("*.part.xlsx"))


def test_export_service_removes_page_limit_and_calendar_omits_unscheduled(tmp_path) -> None:
    start = datetime(2026, 9, 15, tzinfo=UTC)
    scheduled = _task(
        task_id=1,
        title="일정 업무",
        all_day=True,
        starts_at=start,
        ends_at=start + timedelta(days=1),
    )
    unscheduled = replace(scheduled, id=2, title="일정 없음", starts_at=None, ends_at=None)

    class TaskServiceStub:
        def __init__(self) -> None:
            self.queries: list[TaskQuery] = []

        def query(self, query: TaskQuery) -> TaskPage:
            self.queries.append(query)
            return TaskPage((scheduled, unscheduled), 2, query.offset, query.limit)

    class ExporterStub:
        def __init__(self) -> None:
            self.tasks: tuple[Task, ...] = ()

        def export(self, tasks, destination, *, cancel_requested=None):
            self.tasks = tasks
            return destination

    task_service = TaskServiceStub()
    excel = ExporterStub()
    calendar = ExporterStub()
    service = ExportService(task_service, excel, calendar)  # type: ignore[arg-type]

    service.export_excel(TaskQuery(search="일정", limit=1), tmp_path / "tasks.xlsx")
    service.export_calendar(tmp_path / "schedule.ics")

    assert task_service.queries[0].search == "일정"
    assert task_service.queries[0].limit is None
    assert excel.tasks == (scheduled, unscheduled)
    assert calendar.tasks == (scheduled,)
